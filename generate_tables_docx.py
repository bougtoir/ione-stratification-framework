"""Generate a separate editable .docx containing the main tables."""

import os
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SUMMARY_DIR = os.path.join(os.path.dirname(__file__), 'results', 'summary')
OUT_DIR = os.path.join(os.path.dirname(__file__), 'results', 'manuscript')
os.makedirs(OUT_DIR, exist_ok=True)


def _fmt(x, decimals=3):
    if pd.isna(x):
        return '—'
    return f'{x:.{decimals}f}'


def _add_table(doc, df, cols, caption, numeric_cols):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    run.bold = True
    rows = [cols]
    for _, row in df.iterrows():
        r = []
        for c in cols:
            src = numeric_cols[c]
            if src == 'method':
                r.append(row['method'])
            elif src == 'dataset':
                r.append(row['dataset'])
            elif src == 'n_strata':
                r.append(str(int(row['n_strata'])))
            else:
                r.append(_fmt(row.get(src, np.nan), 3 if 'ARI' in src or 'C1' in src or 'W_' in src or 'bias' in src else 2))
        rows.append(r)
    table = doc.add_table(rows=len(rows), cols=len(cols))
    table.style = 'Table Grid'
    for i, row in enumerate(rows):
        cells = table.rows[i].cells
        for j, val in enumerate(row):
            cells[j].text = str(val)
            if i == 0:
                for paragraph in cells[j].paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    doc.add_paragraph()


def main():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)

    phase1 = pd.read_csv(os.path.join(SUMMARY_DIR, 'phase1_primary_summary.csv'))
    _add_table(
        doc, phase1,
        ['Method', 'ARI', 'C1', 'W_true', 'W_est', 'Abs bias crude', 'Abs bias strat', 'Bias reduction', 'Relative reduction'],
        'Table 1. Phase 1 primary scenario (n=2000, K=5): means over simulations.',
        {
            'Method': 'method', 'ARI': 'ARI_mean', 'C1': 'C1_heterogeneity_mean',
            'W_true': 'W_true_mean', 'W_est': 'W_est_mean',
            'Abs bias crude': 'abs_bias_crude_mean', 'Abs bias strat': 'abs_bias_stratified_mean',
            'Bias reduction': 'bias_reduction_mean', 'Relative reduction': 'bias_reduction_relative_mean'
        }
    )

    real = pd.read_csv(os.path.join(SUMMARY_DIR, 'real_data_summary.csv'))
    rd_best = real.loc[real.groupby('dataset')['ARI_mean'].idxmax()].copy()
    _add_table(
        doc, rd_best,
        ['Dataset', 'Method', 'K', 'ARI', 'C1', 'W_est', 'Bias reduction'],
        'Table 2. Best real-data illustration result per dataset.',
        {
            'Dataset': 'dataset', 'Method': 'method', 'K': 'n_strata',
            'ARI': 'ARI_mean', 'C1': 'C1_heterogeneity_mean', 'W_est': 'W_est_mean',
            'Bias reduction': 'bias_reduction_mean'
        }
    )

    out = os.path.join(OUT_DIR, 'tables_separate.docx')
    doc.save(out)
    print(f'[generate_tables_docx] {out}')


if __name__ == '__main__':
    main()
