"""
Minimal markdown-to-docx converter for submission documents.
Supports headings, paragraphs, bold/italic/code, bullet lists, blockquotes and tables.
"""

import re
import sys
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _apply_inline(runs, text):
    """Apply bold/italic/code formatting to a paragraph text and add runs."""
    # Split by bold
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|\`[^`]+\`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = runs.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = runs.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = runs.add_run(part[1:-1])
            run.font.name = 'Courier New'
        else:
            runs.add_run(part)


def convert(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = [l.rstrip('\n') for l in f]

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue
        if line.startswith('## '):
            doc.add_heading(line[3:], level=1)
            i += 1
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:], level=2)
            i += 1
            continue
        if line.startswith('#### '):
            doc.add_heading(line[5:], level=3)
            i += 1
            continue

        # Blockquote
        if line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            _apply_inline(p, line[2:])
            i += 1
            continue

        # Bullet list
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            _apply_inline(p, line[2:])
            i += 1
            continue

        # Numbered list
        m = re.match(r'^(\d+)\.\s+', line)
        if m:
            p = doc.add_paragraph(style='List Number')
            _apply_inline(p, line[m.end():])
            i += 1
            continue

        # Table
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].split('|')]
                # Remove leading/trailing empties
                if cells and cells[0] == '':
                    cells = cells[1:]
                if cells and cells[-1] == '':
                    cells = cells[:-1]
                # Skip separator rows
                if not all(re.match(r':?-+:?', c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                n_cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=n_cols)
                table.style = 'Table Grid'
                for r_idx, r in enumerate(rows):
                    for c_idx in range(n_cols):
                        cell = table.rows[r_idx].cells[c_idx]
                        if c_idx < len(r):
                            _apply_inline(cell.paragraphs[0], r[c_idx])
                        if r_idx == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            continue

        # Plain paragraph
        p = doc.add_paragraph()
        _apply_inline(p, line)
        i += 1

    doc.save(docx_path)
    print(f'[md_to_docx] {md_path} -> {docx_path}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print('Usage: python3 md_to_docx.py input.md output.docx')
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
