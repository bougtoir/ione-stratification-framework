"""
Markdown-to-docx converter for RSM manuscripts.

Supports:
- Headings (#, ##, ###, ####)
- Paragraphs with **bold**, *italic*, `code`
- Bullet and numbered lists
- Blockquotes
- Tables (pipe syntax)
- Vancouver numbered citations [key] -> [n]
- Figures with alt-text
- Page breaks via a {{PAGE}} marker line
- Word OMML equations for mathematical expressions
"""

import re
import os
import copy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from collections import OrderedDict


class CiteManager:
    """Vancouver-style numbered citation manager."""
    def __init__(self):
        self._refs = OrderedDict()
        self._cited = []
        self._key_to_num = {}

    def register(self, key, author_short, year, full_ref, doi=None):
        self._refs[key] = {'author': author_short, 'year': year, 'full': full_ref, 'doi': doi}
        return key

    def cite(self, key, narrative=False):
        if key not in self._refs:
            raise KeyError(f'Citation key {key} not registered')
        if key not in self._key_to_num:
            self._cited.append(key)
            self._key_to_num[key] = len(self._cited)
        n = self._key_to_num[key]
        r = self._refs[key]
        if narrative:
            return f"{r['author']} [{n}]"
        return f'[{n}]'

    def write_reference_list(self, doc):
        if not self._cited:
            return
        doc.add_heading('References', level=1)
        for i, key in enumerate(self._cited, 1):
            p = doc.add_paragraph()
            ref = self._refs[key]
            text = f'{i}. {ref["full"]}'
            doi = ref.get('doi')
            if doi:
                text += f' doi: {doi}'
            p.add_run(text)


def _apply_inline(runs, text, cite_manager=None):
    """Apply inline formatting and citation replacement to a paragraph or cell."""
    if cite_manager:
        text = re.sub(
            r'\[([a-zA-Z0-9_]+)(?:\|([^\]]+))?\]',
            lambda m: cite_manager.cite(m.group(1), narrative=(m.group(2) == 'narrative')),
            text
        )
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|\`[^`]+\`|\[\d+\])', text)
    prev_was_cite = False
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = runs.add_run(part[2:-2])
            run.bold = True
            prev_was_cite = False
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = runs.add_run(part[1:-1])
            run.italic = True
            prev_was_cite = False
        elif part.startswith('`') and part.endswith('`'):
            run = runs.add_run(part[1:-1])
            run.font.name = 'Courier New'
            prev_was_cite = False
        elif re.fullmatch(r'\[\d+\]', part):
            n = part[1:-1]
            if prev_was_cite:
                runs.add_run(',')
            run = runs.add_run(n)
            run.font.superscript = True
            prev_was_cite = True
        else:
            runs.add_run(part)
            prev_was_cite = False


def convert(md_path, docx_path, cite_manager=None, figure_dir=None):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = [l.rstrip('\n') for l in f]

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)

    # Force headings to use Times New Roman and black text (journal formatting)
    for lvl in range(0, 4):
        try:
            h_style = doc.styles[f'Heading {lvl}']
        except KeyError:
            continue
        h_style.font.name = 'Times New Roman'
        h_style.font.color.rgb = RGBColor(0, 0, 0)

    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue

        stripped = line.strip()

        # Page break marker
        if stripped == '{{PAGE}}':
            doc.add_page_break()
            i += 1
            continue

        # Reference list marker
        if stripped == '{{REFS}}':
            if cite_manager:
                cite_manager.write_reference_list(doc)
            i += 1
            continue

        # Figure: ![caption](path)
        m = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', stripped)
        if m:
            caption, path = m.group(1), m.group(2)
            if figure_dir:
                path = os.path.join(figure_dir, path)
            if os.path.exists(path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(6)
                pic = p.add_run().add_picture(path, width=Inches(5.8))
                try:
                    if pic._inline is not None and pic._inline.docPr is not None:
                        pic._inline.docPr.set('descr', caption)
                except Exception:
                    pass
            else:
                doc.add_paragraph(f'[Figure not found: {path}]')
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_before = Pt(12)
            cap_run = cap.add_run(caption)
            cap_run.bold = True
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue
        if line.startswith('## '):
            doc.add_heading(line[3:], level=1)
            i += 1
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:], level=2)
            i += 1
            continue
        if line.startswith('#### '):
            doc.add_heading(line[5:], level=3)
            i += 1
            continue

        # Blockquote
        if line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            _apply_inline(p, line[2:], cite_manager)
            i += 1
            continue

        # Bullet list
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            _apply_inline(p, line[2:], cite_manager)
            i += 1
            continue

        # Numbered list
        m = re.match(r'^(\d+)\.\s+', line)
        if m:
            p = doc.add_paragraph(style='List Number')
            _apply_inline(p, line[m.end():], cite_manager)
            i += 1
            continue

        # Table
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].split('|')]
                if cells and cells[0] == '':
                    cells = cells[1:]
                if cells and cells[-1] == '':
                    cells = cells[:-1]
                # Skip separator rows
                if not all(re.match(r':?-+:?', c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                n_cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=n_cols)
                table.style = 'Table Grid'
                for r_idx, r in enumerate(rows):
                    for c_idx in range(n_cols):
                        cell = table.rows[r_idx].cells[c_idx]
                        if c_idx < len(r):
                            _apply_inline(cell.paragraphs[0], r[c_idx], cite_manager)
                        if r_idx == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            continue

        # Plain paragraph
        p = doc.add_paragraph()
        _apply_inline(p, line, cite_manager)
        if stripped.startswith('*') and ('Table' in stripped or 'Figure' in stripped):
            p.paragraph_format.space_before = Pt(12)
        i += 1

    apply_math_to_doc(doc)
    doc.save(docx_path)
    print(f'[md_to_rsm_docx] {md_path} -> {docx_path}')


# ---------------------------------------------------------------------------
# Word equation (OMML) insertion for mathematical expressions
# ---------------------------------------------------------------------------

_MATH_PATTERNS = [
    (r'C1_excess', 'c1excess'),
    (r'W_(true|est)_excess', 'wexcess'),
    (r'W_excess', 'wexcess2'),
    (r'eta\^2', 'etasq'),
    (r'theta_k', 'thetasub'),
    (r'tau_k', 'tausub'),
    (r'tau\(X_i\)', 'tauxi'),
    (r'v_w', 'vsub'),
    (r'n_k', 'nsub'),
    (r'P\(Y=1\|A=1\) - P\(Y=1\|A=0\)', 'riskdiff'),
    (r'C1 = 1 - I\^2 = 1 - tau\^2 / \(tau\^2 \+ v_w\)', 'c1eq'),
    (r'W = sum_k n_k \(tau_k - tau\)\^2 / sum_i \(tau_i - tau\)\^2', 'weq'),
    (r'C1 = 1 - I\^2', 'c1def'),
    (r'Y ~ X \+ A \+ X\*A', 'ymodel'),
    (r'1 - \|bias_([a-z_]+)\| / \|bias_([a-z_]+)\|', 'frac'),
    (r'\|bias_([a-z_]+)\| - \|bias_([a-z_]+)\|', 'diff'),
    (r'\|Y - p\^\|', 'absphat'),
    (r'tau\^2', 'tausq'),
    (r'W_(true|est)', 'wsub'),
    (r'K_(true|est)', 'ksub'),
    (r'I\^2', 'isq'),
    (r'(?<![A-Za-z0-9])p\^', 'phat'),
]


def _math_text(oMath, text):
    r = OxmlElement('m:r')
    t = OxmlElement('m:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    oMath.append(r)


def _math_superscript(oMath, base, sup):
    ss = OxmlElement('m:sSup')
    e = OxmlElement('m:e')
    _math_text(e, base)
    sup_el = OxmlElement('m:sup')
    _math_text(sup_el, sup)
    ss.append(e)
    ss.append(sup_el)
    oMath.append(ss)


def _math_subscript(oMath, base, sub):
    ss = OxmlElement('m:sSub')
    e = OxmlElement('m:e')
    _math_text(e, base)
    sub_el = OxmlElement('m:sub')
    _math_text(sub_el, sub)
    ss.append(e)
    ss.append(sub_el)
    oMath.append(ss)


def _math_fraction(oMath, num, den):
    f = OxmlElement('m:f')
    num_el = OxmlElement('m:num')
    _math_text(num_el, num)
    den_el = OxmlElement('m:den')
    _math_text(den_el, den)
    f.append(num_el)
    f.append(den_el)
    oMath.append(f)


def _math_accent(oMath, base, accent='^'):
    acc = OxmlElement('m:acc')
    accPr = OxmlElement('m:accPr')
    chr = OxmlElement('m:chr')
    chr.set(qn('m:val'), accent)
    accPr.append(chr)
    e = OxmlElement('m:e')
    _math_text(e, base)
    acc.append(accPr)
    acc.append(e)
    oMath.append(acc)


def _make_math_omath(kind, match):
    oMath = OxmlElement('m:oMath')
    if kind == 'riskdiff':
        _math_text(oMath, 'P(Y=1|A=1) - P(Y=1|A=0)')
    elif kind == 'c1eq':
        _math_text(oMath, 'C1 = 1 - ')
        _math_superscript(oMath, 'I', '2')
        _math_text(oMath, ' = 1 - ')
        f = OxmlElement('m:f')
        num = OxmlElement('m:num')
        _math_superscript(num, 'τ', '2')
        den = OxmlElement('m:den')
        _math_superscript(den, 'τ', '2')
        _math_text(den, ' + ')
        _math_subscript(den, 'v', 'w')
        f.append(num)
        f.append(den)
        oMath.append(f)
    elif kind == 'weq':
        _math_text(oMath, 'W = ')
        f = OxmlElement('m:f')
        num = OxmlElement('m:num')
        _math_subscript(num, 'Σ', 'k')
        _math_text(num, ' ')
        _math_subscript(num, 'n', 'k')
        _math_text(num, ' ')
        group = OxmlElement('m:e')
        _math_text(group, '(')
        _math_subscript(group, 'τ', 'k')
        _math_text(group, ' - ')
        _math_text(group, 'τ')
        _math_text(group, ')')
        ssup = OxmlElement('m:sSup')
        ssup.append(group)
        sup2 = OxmlElement('m:sup')
        _math_text(sup2, '2')
        ssup.append(sup2)
        num.append(ssup)
        den = OxmlElement('m:den')
        _math_subscript(den, 'Σ', 'i')
        _math_text(den, ' ')
        group2 = OxmlElement('m:e')
        _math_text(group2, '(')
        _math_subscript(group2, 'τ', 'i')
        _math_text(group2, ' - ')
        _math_text(group2, 'τ')
        _math_text(group2, ')')
        ssup2 = OxmlElement('m:sSup')
        ssup2.append(group2)
        sup2b = OxmlElement('m:sup')
        _math_text(sup2b, '2')
        ssup2.append(sup2b)
        den.append(ssup2)
        f.append(num)
        f.append(den)
        oMath.append(f)
    elif kind == 'c1def':
        _math_text(oMath, 'C1 = 1 - ')
        _math_superscript(oMath, 'I', '2')
    elif kind == 'ymodel':
        _math_text(oMath, 'Y ~ X + A + X×A')
    elif kind == 'isq':
        _math_superscript(oMath, 'I', '2')
    elif kind == 'tausq':
        _math_superscript(oMath, 'τ', '2')
    elif kind == 'wsub':
        _math_subscript(oMath, 'W', match.group(1))
    elif kind == 'ksub':
        _math_subscript(oMath, 'K', match.group(1))
    elif kind == 'diff':
        _math_text(oMath, f'|bias_{match.group(1)}| - |bias_{match.group(2)}|')
    elif kind == 'frac':
        _math_text(oMath, '1 - ')
        _math_fraction(oMath, f'|bias_{match.group(1)}|', f'|bias_{match.group(2)}|')
    elif kind == 'phat':
        _math_accent(oMath, 'p', accent='^')
    elif kind == 'absphat':
        _math_text(oMath, '|Y - ')
        _math_accent(oMath, 'p', accent='^')
        _math_text(oMath, '|')
    elif kind == 'thetasub':
        _math_subscript(oMath, 'θ', 'k')
    elif kind == 'tausub':
        _math_subscript(oMath, 'τ', 'k')
    elif kind == 'tauxi':
        _math_text(oMath, 'τ(')
        _math_subscript(oMath, 'X', 'i')
        _math_text(oMath, ')')
    elif kind == 'vsub':
        _math_subscript(oMath, 'v', 'w')
    elif kind == 'nsub':
        _math_subscript(oMath, 'n', 'k')
    elif kind == 'etasq':
        _math_superscript(oMath, 'η', '2')
    elif kind == 'c1excess':
        _math_subscript(oMath, 'C1', 'excess')
    elif kind == 'wexcess':
        _math_subscript(oMath, 'W', match.group(1) + ' excess')
    elif kind == 'wexcess2':
        _math_subscript(oMath, 'W', 'excess')
    return oMath


def _text_run(text, rPr):
    r = OxmlElement('w:r')
    if rPr is not None:
        r.append(copy.deepcopy(rPr))
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    return r


def _split_text_to_elems(text, rPr, patterns):
    # Normalise Unicode superscripts to caret notation so OMML patterns catch them
    text = text.replace('²', '^2')
    best = None
    for pat, kind in patterns:
        m = re.search(pat, text)
        if m:
            if best is None or m.start() < best[0].start():
                best = (m, kind)
    if best is None:
        return [_text_run(text, rPr)]
    m, kind = best
    pre = text[:m.start()]
    post = text[m.end():]
    elems = []
    if pre:
        elems.append(_text_run(pre, rPr))
    elems.append(_make_math_omath(kind, m))
    if post:
        elems.extend(_split_text_to_elems(post, rPr, patterns))
    return elems


def _split_run(run, patterns):
    text = run.text
    if not text:
        return None
    rPr = run._r.rPr
    for pat, kind in patterns:
        if re.search(pat, text):
            return _split_text_to_elems(text, rPr, patterns)
    return None


def apply_math_to_doc(doc):
    """Replace plain-text mathematical expressions with Word OMML equations."""
    paras = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)
    for p in paras:
        for run in list(p.runs):
            new_elems = _split_run(run, _MATH_PATTERNS)
            if new_elems:
                for el in new_elems:
                    run._r.addprevious(el)
                p._p.remove(run._r)
