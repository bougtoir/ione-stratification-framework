"""Generate a separate editable .docx containing all RSM tables."""

import os
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt

SUMMARY_DIR = os.path.join(os.path.dirname(__file__), 'results', 'summary')
OUT_DIR = os.path.join(os.path.dirname(__file__), 'results', 'manuscript', 'rsm_submission')
os.makedirs(OUT_DIR, exist_ok=True)


def _fmt(x, decimals=3):
    if pd.isna(x):
        return '—'
    return f'{x:.{decimals}f}'


def _read(name):
    path = os.path.join(SUMMARY_DIR, name)
    if not os.path.exists(path):
        return None
    return pd.read_csv(path)


def _add_table(doc, df, cols, caption, numeric_cols, decimal_overrides=None):
    """Add a table to the docx. numeric_cols maps output column -> source column."""
    if df is None or df.empty:
        return
    p = doc.add_paragraph()
    p.alignment = 1  # CENTER
    run = p.add_run(caption)
    run.bold = True
    rows = [cols]
    overrides = decimal_overrides or {}
    for _, row in df.iterrows():
        r = []
        for c in cols:
            src = numeric_cols[c]
            if src in ('method', 'dataset', 'condition'):
                r.append(str(row.get(src, '')))
            elif src == 'n_strata':
                v = row.get(src, np.nan)
                r.append(str(int(v)) if not pd.isna(v) else '—')
            else:
                dec = overrides.get(c, 3 if ('ARI' in src or 'C1' in src or 'W_' in src or 'bias' in src or 'I2' in src) else 2)
                r.append(_fmt(row.get(src, np.nan), dec))
        rows.append(r)
    table = doc.add_table(rows=len(rows), cols=len(cols))
    table.style = 'Table Grid'
    for i, row in enumerate(rows):
        cells = table.rows[i].cells
        for j, val in enumerate(row):
            cells[j].text = str(val)
            if i == 0:
                for paragraph in cells[j].paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    doc.add_paragraph()


def _sens_row(df, method, n, z, zx, k, metric):
    if df is None or df.empty or metric not in df.columns:
        return np.nan
    sub = df[(df['method'] == method) & (df['n'] == n) &
             (df['z_effect_scale'] == z) & (df['zx_influence_scale'] == zx) &
             (df['n_strata'] == k)]
    return sub[metric].iloc[0] if not sub.empty else np.nan


def _sens_df(df, methods, conditions):
    """Build a DataFrame for sensitivity tables with Monte Carlo SEs."""
    rows = []
    for cond in conditions:
        for method in methods:
            n, z, zx, k = cond.get('n', 2000), cond.get('z', 1.0), cond.get('zx', 1.0), cond.get('k', 5)
            re_bias = _sens_row(df, method, n, z, zx, k, 'abs_bias_re_mean')
            if pd.isna(re_bias):
                continue
            rows.append({
                'condition': cond.get('label', ''),
                'method': method,
                'ARI_mean': _sens_row(df, method, n, z, zx, k, 'ARI_mean'),
                'ARI_se': _sens_row(df, method, n, z, zx, k, 'ARI_se'),
                'C1_heterogeneity_mean': _sens_row(df, method, n, z, zx, k, 'C1_heterogeneity_mean'),
                'C1_heterogeneity_se': _sens_row(df, method, n, z, zx, k, 'C1_heterogeneity_se'),
                'W_est_mean': _sens_row(df, method, n, z, zx, k, 'W_est_mean'),
                'W_est_se': _sens_row(df, method, n, z, zx, k, 'W_est_se'),
                'abs_bias_re_mean': re_bias,
                'abs_bias_re_se': _sens_row(df, method, n, z, zx, k, 'abs_bias_re_se'),
                'bias_reduction_relative_re_mean': _sens_row(df, method, n, z, zx, k, 'bias_reduction_relative_re_mean'),
                'bias_reduction_relative_re_se': _sens_row(df, method, n, z, zx, k, 'bias_reduction_relative_re_se'),
                're_I2_mean': _sens_row(df, method, n, z, zx, k, 're_I2_mean'),
            })
    return pd.DataFrame(rows)


def main():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)

    primary = _read('rsm_ipd_primary_summary.csv')
    full = _read('rsm_ipd_full_summary.csv')
    real = _read('real_data_summary.csv')
    sensitivity = _read('rsm_ipd_sensitivity_full_summary.csv')
    nonlinearity = _read('rsm_ipd_nonlinearity_full_summary.csv')

    # Determine top methods from primary scenario
    top_methods = ['1B_residual', 'PS_propensity_score', 'GMM', 'Prognostic_score', '2B_clustering']
    if primary is not None and not primary.empty:
        best = primary.loc[primary['bias_reduction_relative_re_mean'].idxmax()]['method']
        top_methods = list(dict.fromkeys([best] + top_methods))

    # Table 1: Primary scenario
    if primary is not None and not primary.empty:
        _add_table(
            doc, primary,
            ['Method', 'ARI', 'C1', 'W_true', 'W_est', 'Crude bias', 'Stratified bias',
             'RE bias', 'Relative reduction strat', 'Relative reduction RE', 'RE I2'],
            'Table 1. Primary IPD scenario (n=2000, 10 studies, K=5): means over 50 simulations.',
            {
                'Method': 'method', 'ARI': 'ARI_mean', 'C1': 'C1_heterogeneity_mean',
                'W_true': 'W_true_mean', 'W_est': 'W_est_mean',
                'Crude bias': 'abs_bias_crude_mean', 'Stratified bias': 'abs_bias_stratified_mean',
                'RE bias': 'abs_bias_re_mean',
                'Relative reduction strat': 'bias_reduction_relative_mean',
                'Relative reduction RE': 'bias_reduction_relative_re_mean',
                'RE I2': 're_I2_mean'
            },
            decimal_overrides={'Crude bias': 5, 'Stratified bias': 5, 'RE bias': 5,
                               'Relative reduction strat': 3, 'Relative reduction RE': 3}
        )

    # Table 2: Strata count sensitivity
    if full is not None and not full.empty:
        _add_table(
            doc, full[full['method'].isin(top_methods)],
            ['K', 'Method', 'ARI', 'RE bias', 'Relative reduction RE'],
            'Table 2. Sensitivity of random-effects ATE bias reduction to the number of strata.',
            {
                'K': 'n_strata', 'Method': 'method', 'ARI': 'ARI_mean',
                'RE bias': 'abs_bias_re_mean', 'Relative reduction RE': 'bias_reduction_relative_re_mean'
            },
            decimal_overrides={'RE bias': 5, 'Relative reduction RE': 3}
        )

    # Table 3: Real data illustration
    if real is not None and not real.empty:
        real_non_oracle = real[~real['method'].str.startswith('baseline_')]
        rd_best = real_non_oracle.loc[real_non_oracle.groupby('dataset')['ARI_mean'].idxmax()].copy()
        _add_table(
            doc, rd_best,
            ['Dataset', 'Method', 'K', 'ARI', 'C1', 'W_est', 'Bias reduction'],
            'Table 3. Best real-data illustration result per dataset (oracle baselines excluded).',
            {
                'Dataset': 'dataset', 'Method': 'method', 'K': 'n_strata',
                'ARI': 'ARI_mean', 'C1': 'C1_heterogeneity_mean', 'W_est': 'W_est_mean',
                'Bias reduction': 'bias_reduction_mean'
            }
        )

    # Tables 4-6: sensitivity scenarios
    sample_df = _sens_df(sensitivity, top_methods, [
        {'label': 'n=500', 'n': 500, 'z': 1.0, 'zx': 1.0, 'k': 5},
        {'label': 'n=2000', 'n': 2000, 'z': 1.0, 'zx': 1.0, 'k': 5},
        {'label': 'n=10000', 'n': 10000, 'z': 1.0, 'zx': 1.0, 'k': 5},
    ])
    if not sample_df.empty:
        _add_table(
            doc, sample_df,
            ['Condition', 'Method', 'ARI', 'ARI SE', 'C1', 'C1 SE', 'W_est', 'W_est SE', 'RE bias', 'RE bias SE',
             'Relative reduction RE', 'Relative reduction RE SE', 'RE I2'],
            'Table 4. Sample-size sensitivity (K=5, z=1.0, zx=1.0): means over 10 simulations.',
            {
                'Condition': 'condition', 'Method': 'method', 'ARI': 'ARI_mean', 'ARI SE': 'ARI_se',
                'C1': 'C1_heterogeneity_mean', 'C1 SE': 'C1_heterogeneity_se',
                'W_est': 'W_est_mean', 'W_est SE': 'W_est_se',
                'RE bias': 'abs_bias_re_mean', 'RE bias SE': 'abs_bias_re_se',
                'Relative reduction RE': 'bias_reduction_relative_re_mean',
                'Relative reduction RE SE': 'bias_reduction_relative_re_se',
                'RE I2': 're_I2_mean'
            },
            decimal_overrides={'RE bias': 5, 'RE bias SE': 5, 'Relative reduction RE': 3, 'Relative reduction RE SE': 3}
        )

    zx_df = _sens_df(sensitivity, top_methods, [
        {'label': 'zx=0.2', 'n': 2000, 'z': 1.0, 'zx': 0.2, 'k': 5},
        {'label': 'zx=0.5', 'n': 2000, 'z': 1.0, 'zx': 0.5, 'k': 5},
        {'label': 'zx=1.0', 'n': 2000, 'z': 1.0, 'zx': 1.0, 'k': 5},
    ])
    if not zx_df.empty:
        _add_table(
            doc, zx_df,
            ['Condition', 'Method', 'ARI', 'ARI SE', 'C1', 'C1 SE', 'W_est', 'W_est SE', 'RE bias', 'RE bias SE',
             'Relative reduction RE', 'Relative reduction RE SE', 'RE I2'],
            'Table 5. Sensitivity to Z-to-X influence strength (n=2000, K=5, z=1.0, zx=0.2, 0.5, 1.0): means over 10 simulations.',
            {
                'Condition': 'condition', 'Method': 'method', 'ARI': 'ARI_mean', 'ARI SE': 'ARI_se',
                'C1': 'C1_heterogeneity_mean', 'C1 SE': 'C1_heterogeneity_se',
                'W_est': 'W_est_mean', 'W_est SE': 'W_est_se',
                'RE bias': 'abs_bias_re_mean', 'RE bias SE': 'abs_bias_re_se',
                'Relative reduction RE': 'bias_reduction_relative_re_mean',
                'Relative reduction RE SE': 'bias_reduction_relative_re_se',
                'RE I2': 're_I2_mean'
            },
            decimal_overrides={'RE bias': 5, 'RE bias SE': 5, 'Relative reduction RE': 3, 'Relative reduction RE SE': 3}
        )

    z_df = _sens_df(sensitivity, top_methods, [
        {'label': 'z=0.5', 'n': 2000, 'z': 0.5, 'zx': 1.0, 'k': 5},
        {'label': 'z=1.0', 'n': 2000, 'z': 1.0, 'zx': 1.0, 'k': 5},
        {'label': 'z=2.0', 'n': 2000, 'z': 2.0, 'zx': 1.0, 'k': 5},
    ])
    if not z_df.empty:
        _add_table(
            doc, z_df,
            ['Condition', 'Method', 'ARI', 'ARI SE', 'C1', 'C1 SE', 'W_est', 'W_est SE', 'RE bias', 'RE bias SE',
             'Relative reduction RE', 'Relative reduction RE SE', 'RE I2'],
            'Table 6. Sensitivity to Z-to-Y effect strength (n=2000, K=5, zx=1.0, z=0.5, 1.0, 2.0): means over 10 simulations.',
            {
                'Condition': 'condition', 'Method': 'method', 'ARI': 'ARI_mean', 'ARI SE': 'ARI_se',
                'C1': 'C1_heterogeneity_mean', 'C1 SE': 'C1_heterogeneity_se',
                'W_est': 'W_est_mean', 'W_est SE': 'W_est_se',
                'RE bias': 'abs_bias_re_mean', 'RE bias SE': 'abs_bias_re_se',
                'Relative reduction RE': 'bias_reduction_relative_re_mean',
                'Relative reduction RE SE': 'bias_reduction_relative_re_se',
                'RE I2': 're_I2_mean'
            },
            decimal_overrides={'RE bias': 5, 'RE bias SE': 5, 'Relative reduction RE': 3, 'Relative reduction RE SE': 3}
        )

    # Table 7: Nonlinearity comparison
    if nonlinearity is not None and not nonlinearity.empty:
        rows = []
        for method in top_methods:
            n, z, zx, k = 2000, 1.0, 1.0, 5
            lin_bias = _sens_row(sensitivity, method, n, z, zx, k, 'abs_bias_re_mean')
            lin_bias_se = _sens_row(sensitivity, method, n, z, zx, k, 'abs_bias_re_se')
            lin_rel = _sens_row(sensitivity, method, n, z, zx, k, 'bias_reduction_relative_re_mean')
            lin_rel_se = _sens_row(sensitivity, method, n, z, zx, k, 'bias_reduction_relative_re_se')
            non_bias = _sens_row(nonlinearity, method, n, z, zx, k, 'abs_bias_re_mean')
            non_bias_se = _sens_row(nonlinearity, method, n, z, zx, k, 'abs_bias_re_se')
            non_rel = _sens_row(nonlinearity, method, n, z, zx, k, 'bias_reduction_relative_re_mean')
            non_rel_se = _sens_row(nonlinearity, method, n, z, zx, k, 'bias_reduction_relative_re_se')
            if not pd.isna(non_bias):
                rows.append({
                    'method': method,
                    'linear_abs_bias_re_mean': lin_bias,
                    'linear_abs_bias_re_se': lin_bias_se,
                    'linear_bias_reduction_relative_re_mean': lin_rel,
                    'linear_bias_reduction_relative_re_se': lin_rel_se,
                    'nonlinear_abs_bias_re_mean': non_bias,
                    'nonlinear_abs_bias_re_se': non_bias_se,
                    'nonlinear_bias_reduction_relative_re_mean': non_rel,
                    'nonlinear_bias_reduction_relative_re_se': non_rel_se,
                })
        if rows:
            nonlin_df = pd.DataFrame(rows)
            _add_table(
                doc, nonlin_df,
                ['Method', 'Linear RE bias', 'Linear RE bias SE', 'Linear rel reduction', 'Linear rel reduction SE',
                 'Non-linear RE bias', 'Non-linear RE bias SE', 'Non-linear rel reduction', 'Non-linear rel reduction SE'],
                'Table 7. Linear versus non-linear Z-to-X mapping (n=2000, K=5, z=1.0, zx=1.0).',
                {
                    'Method': 'method',
                    'Linear RE bias': 'linear_abs_bias_re_mean',
                    'Linear RE bias SE': 'linear_abs_bias_re_se',
                    'Linear rel reduction': 'linear_bias_reduction_relative_re_mean',
                    'Linear rel reduction SE': 'linear_bias_reduction_relative_re_se',
                    'Non-linear RE bias': 'nonlinear_abs_bias_re_mean',
                    'Non-linear RE bias SE': 'nonlinear_abs_bias_re_se',
                    'Non-linear rel reduction': 'nonlinear_bias_reduction_relative_re_mean',
                    'Non-linear rel reduction SE': 'nonlinear_bias_reduction_relative_re_se'
                },
                decimal_overrides={'Linear RE bias': 5, 'Linear RE bias SE': 5,
                                   'Non-linear RE bias': 5, 'Non-linear RE bias SE': 5,
                                   'Linear rel reduction': 3, 'Linear rel reduction SE': 3,
                                   'Non-linear rel reduction': 3, 'Non-linear rel reduction SE': 3}
            )

    out = os.path.join(OUT_DIR, 'rsm_tables_separate.docx')
    doc.save(out)
    print(f'[generate_rsm_tables] {out}')


if __name__ == '__main__':
    main()
